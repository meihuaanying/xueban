"""题库导入服务：JSON/CSV/XLSX → 校验（pydantic）→ 入库（幂等）。"""

from __future__ import annotations

import contextlib
import csv
import json
import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from openpyxl import load_workbook
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    KnowledgePoint,
    MistakeTag,
    Question,
    QuestionKnowledgePoint,
    QuestionMistakeTag,
    QuestionStatus,
    QuestionType,
)


class QuestionRecord(BaseModel):
    """导入记录的校验契约。"""

    subject: str = Field(min_length=1, max_length=32)
    stage: str = Field(default="junior", max_length=20)
    qtype: Literal["choice", "fill", "short_answer", "essay", "programming"] = "choice"
    stem: str = Field(min_length=5)
    options: dict[str, str] | None = None
    answer: str = Field(min_length=1)
    analysis: str | None = None
    difficulty: int = Field(default=3, ge=1, le=5)
    knowledge_points: list[str] = Field(default_factory=list)
    mistake_tags: list[str] = Field(default_factory=list)
    source: str = Field(default="self-built", max_length=128)


@dataclass(slots=True)
class ImportIssue:
    """逐行问题。"""

    row: int
    message: str


@dataclass(slots=True)
class ImportReport:
    """导入报告。"""

    total: int = 0
    imported: int = 0
    skipped: int = 0
    invalid: int = 0
    issues: list[ImportIssue] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        """是否无非法记录。"""
        return self.invalid == 0


def _split_list(value: Any) -> list[str]:
    """字符串/列表 → 字符串列表。"""
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    if not text:
        return []
    if text.startswith("["):
        try:
            parsed = json.loads(text)
        except ValueError:
            parsed = None
        if isinstance(parsed, list):
            return [str(item).strip() for item in parsed if str(item).strip()]
    return [part.strip() for part in re.split(r"[|,，;；\s]+", text) if part.strip()]


def _split_options(value: Any) -> dict[str, str] | None:
    """字符串/字典 → 选项字典。"""
    if value is None or value == "":
        return None
    if isinstance(value, dict):
        return {str(key): str(item) for key, item in value.items()}
    text = str(value).strip()
    if text.startswith("{"):
        try:
            parsed = json.loads(text)
        except ValueError:
            parsed = None
        if isinstance(parsed, dict):
            return {str(key): str(item) for key, item in parsed.items()}
    options: dict[str, str] = {}
    for part in re.split(r"[|;；]", text):
        pieces = re.split(r"[:：]", part, maxsplit=1)
        if len(pieces) == 2 and pieces[0].strip():
            options[pieces[0].strip()] = pieces[1].strip()
    return options or None


def _normalize(raw: dict[str, Any]) -> dict[str, Any]:
    """把表格/CSV 值归一化为契约所需类型。"""
    data = dict(raw)
    data["options"] = _split_options(data.get("options"))
    data["knowledge_points"] = _split_list(data.get("knowledge_points"))
    data["mistake_tags"] = _split_list(data.get("mistake_tags"))
    if data.get("analysis") in (None, ""):
        data.pop("analysis", None)
    if data.get("source") in (None, ""):
        data.pop("source", None)
    if data.get("stage") in (None, ""):
        data.pop("stage", None)
    difficulty = data.get("difficulty")
    if difficulty not in (None, ""):
        with contextlib.suppress(TypeError, ValueError):
            data["difficulty"] = int(difficulty)
    return data


def load_records(path: Path) -> list[tuple[int, dict[str, Any]]]:
    """加载导入文件，返回 (行号, 原始记录) 列表。"""
    suffix = path.suffix.lower()
    if suffix == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        records = payload.get("questions") if isinstance(payload, dict) else payload
        if not isinstance(records, list):
            raise ValueError("JSON 顶层应为数组或 {questions: [...]}")
        rows: list[tuple[int, dict[str, Any]]] = []
        for index, item in enumerate(records, start=1):
            if isinstance(item, dict):
                rows.append((index, dict(item)))
        return rows
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            return [(index, dict(row)) for index, row in enumerate(reader, start=2)]
    if suffix in (".xlsx", ".xlsm"):
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(item).strip() if item is not None else "" for item in rows[0]]
        collected: list[tuple[int, dict[str, Any]]] = []
        for index, row in enumerate(rows[1:], start=2):
            record = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
            collected.append((index, record))
        return collected
    raise ValueError(f"不支持的文件类型：{suffix}（支持 .json/.csv/.xlsx）")


def _summarize(exc: ValidationError) -> str:
    """压缩校验错误为一行。"""
    first = exc.errors()[0] if exc.errors() else None
    if first is None:
        return "记录格式不正确"
    location = ".".join(str(part) for part in first.get("loc", [])) or "record"
    return f"{location}: {first.get('msg', '格式不正确')}"


async def import_questions(
    session: AsyncSession,
    rows: list[tuple[int, dict[str, Any]]],
    *,
    validate_only: bool = False,
    source_default: str | None = None,
) -> ImportReport:
    """逐条校验并导入（非法记录逐条报错，不中断整体）。"""
    report = ImportReport(total=len(rows))

    kp_rows = (await session.execute(select(KnowledgePoint.code, KnowledgePoint.id))).all()
    kp_map = {row[0]: row[1] for row in kp_rows}
    tag_rows = (await session.execute(select(MistakeTag.code, MistakeTag.id))).all()
    tag_map = {row[0]: row[1] for row in tag_rows}
    existing_rows = (await session.execute(select(Question.subject, Question.stem))).all()
    existing = {(row[0], row[1]) for row in existing_rows}
    seen: set[tuple[str, str]] = set()

    batch: list[tuple[QuestionRecord, list[uuid.UUID], list[uuid.UUID]]] = []
    for row_number, raw in rows:
        try:
            record = QuestionRecord.model_validate(_normalize(raw))
        except ValidationError as exc:
            report.invalid += 1
            report.issues.append(ImportIssue(row_number, _summarize(exc)))
            continue
        if record.qtype == QuestionType.CHOICE.value and not record.options:
            report.invalid += 1
            report.issues.append(ImportIssue(row_number, "选择题必须提供 options"))
            continue
        missing_kp = sorted({code for code in record.knowledge_points if code not in kp_map})
        if missing_kp:
            report.invalid += 1
            report.issues.append(
                ImportIssue(row_number, f"未知知识点编码：{', '.join(missing_kp)}")
            )
            continue
        missing_tags = sorted({code for code in record.mistake_tags if code not in tag_map})
        if missing_tags:
            report.invalid += 1
            report.issues.append(
                ImportIssue(row_number, f"未知错因标签：{', '.join(missing_tags)}")
            )
            continue
        key = (record.subject, record.stem)
        if key in existing or key in seen:
            report.skipped += 1
            continue
        seen.add(key)
        batch.append(
            (
                record,
                [kp_map[code] for code in dict.fromkeys(record.knowledge_points)],
                [tag_map[code] for code in dict.fromkeys(record.mistake_tags)],
            )
        )

    if not validate_only:
        for record, kp_ids, tag_ids in batch:
            question = Question(
                subject=record.subject,
                stage=record.stage,
                qtype=QuestionType(record.qtype),
                stem=record.stem,
                options=record.options,
                answer=record.answer,
                analysis=record.analysis or f"参考答案：{record.answer}。",
                difficulty=record.difficulty,
                source=source_default or record.source,
                status=QuestionStatus.PUBLISHED,
            )
            session.add(question)
            await session.flush()
            for kp_id in kp_ids:
                session.add(
                    QuestionKnowledgePoint(question_id=question.id, knowledge_point_id=kp_id)
                )
            for tag_id in tag_ids:
                session.add(QuestionMistakeTag(question_id=question.id, mistake_tag_id=tag_id))
        await session.flush()

    report.imported = len(batch)
    return report
